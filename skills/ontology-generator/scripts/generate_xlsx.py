"""
本体 Schema xlsx 生成辅助脚本

用法：
    from generate_xlsx import OntologySchemaGenerator
    gen = OntologySchemaGenerator()
    gen.add_entity("Equipment（设备）", "升降机设备主体", [
        ("equipment_id", "否", "-", "设备编号"),
        ("name", "否", "-", "设备名称"),
        ("status", "否", "active", "设备状态(active/inactive)"),
    ])
    gen.add_relationship("设备结构关系", "设备与部件的组成关系",
        "Equipment（设备）", "Component（部件）", [
        ("relationship_type", "否", "contains", "关系类型"),
        ("hierarchy_level", "否", "1", "层级"),
    ])
    gen.add_chain("设备结构管理链路", "Equipment（设备） → Component（部件） → SubComponent（子部件）")
    gen.save("output.xlsx")
"""
import sys
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "--break-system-packages", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class OntologySchemaGenerator:
    def __init__(self):
        self.entities = []
        self.relationships = []
        self.chains = []

    def add_entity(self, name: str, description: str, attributes: list):
        """
        添加实体定义。
        attributes: list of (attr_name, nullable, default, desc)
        """
        self.entities.append({
            "name": name,
            "description": description,
            "attributes": attributes
        })

    def add_relationship(self, name: str, description: str,
                         source: str, target: str, attributes: list):
        """
        添加关系定义。
        attributes: list of (attr_name, nullable, default, desc)
        """
        self.relationships.append({
            "name": name,
            "description": description,
            "source": source,
            "target": target,
            "attributes": attributes
        })

    def add_chain(self, title: str, path_description: str):
        self.chains.append({"title": title, "path": path_description})

    def save(self, filepath: str):
        wb = Workbook()

        # --- Sheet 1: 实体与属性定义 ---
        ws1 = wb.active
        ws1.title = "实体与属性定义"
        headers1 = ["实体类型名称", "描述", "属性名称", "是否可为空", "默认值", "属性描述"]
        self._write_header(ws1, headers1)

        row = 2
        for entity in self.entities:
            for i, (attr_name, nullable, default, desc) in enumerate(entity["attributes"]):
                if i == 0:
                    ws1.cell(row=row, column=1, value=entity["name"])
                    ws1.cell(row=row, column=2, value=entity["description"])
                    ws1.cell(row=row, column=1).font = Font(name="Arial", bold=True, size=10)
                    ws1.cell(row=row, column=2).font = Font(name="Arial", size=10)
                ws1.cell(row=row, column=3, value=attr_name)
                ws1.cell(row=row, column=4, value=nullable)
                ws1.cell(row=row, column=5, value=default)
                ws1.cell(row=row, column=6, value=desc)
                for col in range(3, 7):
                    ws1.cell(row=row, column=col).font = Font(name="Arial", size=10)
                row += 1

        ws1.column_dimensions['A'].width = 30
        ws1.column_dimensions['B'].width = 25
        ws1.column_dimensions['C'].width = 25
        ws1.column_dimensions['D'].width = 10
        ws1.column_dimensions['E'].width = 12
        ws1.column_dimensions['F'].width = 50

        # --- Sheet 2: 关系与属性定义 ---
        ws2 = wb.create_sheet("关系与属性定义")
        headers2 = ["关系类型名称", "关系描述", "起点实体", "终点实体",
                     "关系属性名称", "是否可为空", "默认值", "属性描述"]
        self._write_header(ws2, headers2)

        row = 2
        for rel in self.relationships:
            for i, (attr_name, nullable, default, desc) in enumerate(rel["attributes"]):
                if i == 0:
                    ws2.cell(row=row, column=1, value=rel["name"])
                    ws2.cell(row=row, column=2, value=rel["description"])
                    ws2.cell(row=row, column=3, value=rel["source"])
                    ws2.cell(row=row, column=4, value=rel["target"])
                    for col in range(1, 5):
                        ws2.cell(row=row, column=col).font = Font(name="Arial", bold=(col == 1), size=10)
                ws2.cell(row=row, column=5, value=attr_name)
                ws2.cell(row=row, column=6, value=nullable)
                ws2.cell(row=row, column=7, value=default)
                ws2.cell(row=row, column=8, value=desc)
                for col in range(5, 9):
                    ws2.cell(row=row, column=col).font = Font(name="Arial", size=10)
                row += 1

        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 30
        ws2.column_dimensions['C'].width = 30
        ws2.column_dimensions['D'].width = 30
        ws2.column_dimensions['E'].width = 25
        ws2.column_dimensions['F'].width = 10
        ws2.column_dimensions['G'].width = 12
        ws2.column_dimensions['H'].width = 45

        # --- Sheet 3: 实体关系图 ---
        ws3 = wb.create_sheet("实体关系图")
        ws3.cell(row=1, column=1, value="||--o{：一对多关系\n||--||：一对一关系")
        ws3.cell(row=1, column=1).font = Font(name="Arial", size=10)
        ws3.cell(row=1, column=1).alignment = Alignment(wrap_text=True)

        for i, chain in enumerate(self.chains):
            text = f"{self._to_chinese_num(i+1)}、{chain['title']}\n核心路径：{chain['path']}"
            ws3.cell(row=i+3, column=1, value=text)
            ws3.cell(row=i+3, column=1).font = Font(name="Arial", size=10)
            ws3.cell(row=i+3, column=1).alignment = Alignment(wrap_text=True)
            ws3.row_dimensions[i+3].height = 40

        ws3.column_dimensions['A'].width = 100

        wb.save(filepath)
        print(f"本体 Schema 已保存至: {filepath}")

    def _write_header(self, ws, headers):
        header_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{chr(64+len(headers))}1"

    @staticmethod
    def _to_chinese_num(n):
        nums = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五",
                6: "六", 7: "七", 8: "八", 9: "九", 10: "十"}
        return nums.get(n, str(n))


if __name__ == "__main__":
    # 示例：快速测试
    gen = OntologySchemaGenerator()
    gen.add_entity("Equipment（设备）", "设备主体", [
        ("equipment_id", "否", "-", "设备编号"),
        ("name", "否", "-", "设备名称"),
        ("status", "否", "active", "状态(active/inactive)"),
    ])
    gen.add_relationship("设备结构关系", "设备与部件组成",
        "Equipment（设备）", "Component（部件）", [
        ("hierarchy_level", "否", "1", "层级"),
    ])
    gen.add_chain("设备结构管理链路", "Equipment → Component → SubComponent")
    gen.save("/tmp/test_ontology.xlsx")
    print("测试完成")
